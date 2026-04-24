"""Dump paragraphs of a .docx with their style names for verification."""
import sys
from pathlib import Path
from docx import Document

src = Path(sys.argv[1])
doc = Document(str(src))

out = []
for i, p in enumerate(doc.paragraphs):
    style = p.style.name
    text = p.text.strip()
    out.append(f"{i:03d} [{style}] {text[:150]}")

for t in doc.tables:
    out.append(f"--- TABLE ({len(t.rows)} rows x {len(t.columns)} cols) ---")

print("\n".join(out))
